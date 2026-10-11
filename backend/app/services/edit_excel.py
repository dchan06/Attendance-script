from backend.app.services.notion_script import format_absence_data
import holidays
from datetime import date, timedelta

from openpyxl import load_workbook

from enum import IntEnum

def get_details_from_user(): 
    name = input("Full name: ") 
    target_month = input("Month (In this exact format \"Feb-Mar\")): ")
    work_hours = input("What are your work hours: ")
    nric = input("Last 3 digits of your NRIC: ")
    role = input("What is your title/role: ")
    return name, target_month, work_hours, nric, role

def get_holidays(period=""): 
    sg_holidays = holidays.Singapore(years=2026, observed=True)
    return sg_holidays

def load_wb(month): 
    wb = load_workbook("Monthly Timesheet.xlsx")

    wb.sheetnames          # ['Jan-Feb', 'Feb-Mar', 'Mar-Apr']
    ws = wb[month]     # select by name
    return wb, ws

def fill_particulars(self): 
    for row in self.ws.iter_rows(min_row=5, max_row=9): 
        match (row[0].value or "").strip():
            case "Name in NRIC:": 
                row[1].value = self.user
            case "NRIC (last four characters e.g. 123N):": 
                row[1].value = self.nric_last_four
            case "Internship Position held:": 
                row[1].value = self.position
                
def change_month_to_numbers(month: str) -> dict[str, int]: 
    start, end = month.split("-")
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    month_nums = {}
    if start not in months or end not in months: 
        print("Error: Invalid month")
        return
    for i, n in enumerate(months): 
        if n == start: 
            month_nums["start"] = i + 1
        if n == end: 
            month_nums["end"] = i + 1
    return month_nums
                
class Col(IntEnum):
    DATE = 0
    DAY = 1
    ATTENDANCE = 2
    WORKING_HOURS = 3
    OVERTIME_HOURS = 4
    REMARKS = 5
                
def fill_dates_and_attendance(self): 
    i = len(self.dates)
    start_pos = 18
    end_pos = 18 + i -1
    absent_dates = format_absence_data()
    for indx, row in enumerate(self.ws.iter_rows(min_row=start_pos, max_row=end_pos)): 
        date = self.dates[indx]
        row[Col.DATE].value = date
        row[Col.DAY].value = date.strftime("%A")
        if date.isoweekday() > 5: 
            row[Col.ATTENDANCE].value = "Non-Work Day"
        elif date in self.holiday_dates: 
            row[Col.ATTENDANCE].value = "Holiday"
            row[Col.REMARKS].value = self.holiday_obj[date]
        elif date in absent_dates: 
            row[Col.ATTENDANCE].value = "Absent"
        elif date.isoweekday() <= 5:  # Mon starts at 1, Sun = 7
            row[Col.ATTENDANCE].value = "Full Day Work"
            row[Col.WORKING_HOURS].value = "10 a.m. to 7 p.m."
        row[0].offset(row=1).value = ""       #Hardcoded to reset the next cell down to "", so that for months with only 30 days there won't be an empty/ cell with bad data
    
class EditExcel: 
    def __init__(self, month: str, work_hours: str, user: str, nric: str, pos: str):
        self.month=month
        self.work_hours=work_hours
        self.user=user
        self.nric_last_four=nric
        self.position=pos
    def edit_excel(self):
        self.wb, self.ws = load_wb(self.month)
        self.fill_particulars()
        self.dates = []
        month_in_nums = change_month_to_numbers(self.month)
        current = date(2026, month_in_nums["start"], 21)
        end = date(2026, month_in_nums["end"], 20)
        while current <= end: 
            self.dates.append(current)
            current += timedelta(days=1)
        
        self.holiday_obj = get_holidays()   # outputs a dict with key -> date and value -> holiday name (str)
        self.holiday_dates = list(self.holiday_obj.keys())
        
        self.fill_dates()
        self.wb.save("Monthly Timesheet.xlsx")
        
    fill_particulars = fill_particulars
    fill_dates = fill_dates_and_attendance
        
def main(): 
    name, target_month, work_hours, nric, role = get_details_from_user()
    edit = EditExcel(target_month, "1000-1900", name, nric, role)
    edit.edit_excel()
    
if __name__ == "__main__": 
    main()